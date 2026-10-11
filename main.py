from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware
import fitz  # PyMuPDF สำหรับแกะพิกัดเรขาคณิต
from pythainlp.util import normalize as thai_normalize
import uvicorn
import re
import os
from fastapi.responses import FileResponse
from pdf2docx import Converter
import tempfile
import docx

app = FastAPI(title="e-GovDoc Python Backend Engine")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

def extract_and_clean_pdf(file_bytes: bytes) -> str:
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        all_lines = []
        
        # 1. ดึงข้อความพร้อมพิกัด Bounding Box ของแต่ละบรรทัดอย่างละเอียด
        for page_num, page in enumerate(doc):
            blocks = page.get_text("dict").get("blocks", [])
            for b in blocks:
                if b.get("type") == 0:  # Text block
                    for l in b.get("lines", []):
                        bbox = l.get("bbox")  # (x0, y0, x1, y1)
                        line_text = "".join([span.get("text", "") for span in l.get("spans", "")]).strip()
                        if line_text:
                            all_lines.append({
                                "page": page_num,
                                "x0": bbox[0],
                                "y0": bbox[1],
                                "x1": bbox[2],
                                "y1": bbox[3],
                                "text": line_text
                            })
        
        if not all_lines:
            return ""
        
        # เรียงลำดับบรรทัดตามหน้า และพิกัดจากบนลงล่าง ซ้ายไปขวา
        all_lines.sort(key=lambda item: (item["page"], item["y0"], item["x0"]))
        
        # 2. อัลกอริทึมวิเคราะห์ระยะห่างพิกัด (Geometry-based Line Reflow)
        paragraphs = []
        current_para_lines = []
        prev_line = None
        
        for line in all_lines:
            text = line["text"]
            
            if not prev_line:
                current_para_lines.append(text)
                prev_line = line
                continue
            
            # หากเปลี่ยนหน้า ให้ขึ้นย่อหน้าใหม่
            if line["page"] != prev_line["page"]:
                if current_para_lines:
                    paragraphs.append(" ".join(current_para_lines))
                    current_para_lines = []
                current_para_lines.append(text)
                prev_line = line
                continue
            
            # คำนวณระยะห่างแนวตั้ง (Vertical Gap) และความสูงของบรรทัด
            y_gap = line["y0"] - prev_line["y1"]
            line_height = prev_line["y1"] - prev_line["y0"]
            if line_height <= 0:
                line_height = 14.0
            
            # เกณฑ์ตัดสินใจ: ถ้าระยะห่างไม่เกิน 85% ของความสูงบรรทัด แปลว่าเป็นห่อคำในย่อหน้าเดียวกัน
            threshold = line_height * 0.85
            
            if y_gap <= threshold:
                # เชื่อมบรรทัดเข้าด้วยกัน (Soft Wrap)
                last_t = current_para_lines[-1]
                curr_t = text
                
                is_thai_left = bool(re.search(r'[\u0E00-\u0E7F]$', last_t))
                is_thai_right = bool(re.search(r'^[\u0E00-\u0E7F]', curr_t))
                
                if is_thai_left and is_thai_right:
                    current_para_lines[-1] = last_t + curr_t
                else:
                    current_para_lines[-1] = last_t + " " + curr_t
            else:
                # ระยะห่างกว้างกว่าปกติ แปลว่าเป็นขึ้นย่อหน้าใหม่หรือหัวข้อใหม่ (Hard Break)
                if current_para_lines:
                    paragraphs.append(" ".join(current_para_lines))
                current_para_lines = [text]
            
            prev_line = line
            
        if current_para_lines:
            paragraphs.append(" ".join(current_para_lines))
        
        full_text = "\n\n".join(paragraphs)
        
        # 3. 👑 ล็อกกฎสระอำเด็ดขาด 100%: พยัญชนะ [ก-ฮ] + ช่องว่าง (\s+) + สระอา (า) -> สระอำ (ำ)
        full_text = re.sub(r'([ก-ฮ])\s+([า])', r'\1ำ', full_text)
        
        # 4. ทำความสะอาดมาตรฐานด้วย PyThaiNLP
        normalized_text = thai_normalize(full_text)
        
        # 5. จัดระเบียบช่องว่างย่อหน้าให้สะอาดเรียบร้อย
        final_paras = [re.sub(r' {2,}', ' ', p).strip() for p in normalized_text.split("\n\n") if p.strip()]
        
        return "\n\n".join(final_paras)
        
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"ไม่สามารถอ่านไฟล์ PDF ได้: {str(e)}")

@app.get("/")
def root():
    return {"status": "ok", "message": "e-GovDoc Python Backend (Geometry Engine) is running smoothly!"}

@app.post("/api/extract-pdf")
async def extract_pdf_endpoint(file: UploadFile = File(...)):
    file_bytes = await file.read()
    extracted_text = extract_and_clean_pdf(file_bytes)
    return {"success": True, "text": extracted_text}

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 8000))
    uvicorn.run("main:app", host="0.0.0.0", port=port)

@app.post("/api/pdf-to-word")
async def pdf_to_word_endpoint(file: UploadFile = File(...)):
    if not file.filename.lower().endswith(".pdf"):
        raise HTTPException(status_code=400, detail="กรุณาอัปโหลดไฟล์ PDF เท่านั้น")
    
    pdf_fd, pdf_path = tempfile.mkstemp(suffix=".pdf")
    os.close(pdf_fd)
    docx_path = pdf_path.replace(".pdf", ".docx")
    
    try:
        content = await file.read()
        with open(pdf_path, "wb") as f:
            f.write(content)
            
        # 1. แปลง PDF เป็น Word ด้วย pdf2docx ตามปกติ
        cv = Converter(pdf_path)
        cv.convert(docx_path, start=0, end=None)
        cv.close()
        
        # 2. เปิดไฟล์ Word ผลลัพธ์ด้วย python-docx เพื่อซ่อมแซมสระอำเด็ดขาด
        doc = docx.Document(docx_path)
        
        def fix_sara_am(text):
            if not text:
                return text
            # กฎสระอำ: พยัญชนะไทย [ก-ฮ] + ช่องว่าง (\s+) + สระอา (า) -> แปลงเป็นสระอำ (ำ) ทันที
            return re.sub(r'([ก-ฮ])\s+([า])', r'\1ำ', text)

        # แก้ไขในทุกย่อหน้า (Paragraphs)
        for p in doc.paragraphs:
            if p.text:
                fixed_text = fix_sara_am(p.text)
                if fixed_text != p.text:
                    p.text = fixed_text

        # แก้ไขในทุกตาราง (Tables) ถ้ามี
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if p.text:
                            fixed_text = fix_sara_am(p.text)
                            if fixed_text != p.text:
                                p.text = fixed_text
                                
        # บันทึกไฟล์ Word ที่ซ่อมแซมสระอำแล้วทับลงไป
        doc.save(docx_path)
        
        # 3. ส่งไฟล์ Word กลับไปให้ผู้ใช้งานดาวน์โหลด
        output_filename = file.filename.rsplit(".", 1)[0] + ".docx"
        return FileResponse(
            docx_path,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            filename=output_filename
        )
        
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"เกิดข้อผิดพลาดในการแปลงไฟล์: {str(e)}")
    finally:
        if os.path.exists(pdf_path):
            try: os.remove(pdf_path)
            except: pass
